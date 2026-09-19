"""Builds the standard Attur HUD fever / dengue line-list template (no patient data)."""
import os
import shutil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'template', 'Attur_HUD_line_list_template.xlsx')
APP_COPY = None  # the app serves template/ directly
ROWS = 1500  # rows prepared with formats, formulas and validation

TEAL, TEAL_SOFT, INK, GREY = '0E6E5F', 'DCEFE9', '17211E', '75827D'
HEAD_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=TEAL)
AUTO_FILL = PatternFill('solid', fgColor='F1F4F2')
THIN = Side(style='thin', color='C9D1CD')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

BLOCKS = ['Attur', 'Attur Mpty', 'Narasingapuram Mpty', 'Ayothiyapattinam', 'Gangavalli', 'Panamarathupatti',
          'Pethanaickenpalayam', 'Thalaivasal', 'Valapadi', 'Yercaud']
PHCS = ['Attur UPHC', 'Narasingapuram UPHC', 'Malliyakarai', 'Thennankudipalayam', 'Kothampadi', 'Keeripatti', 'Manjini',
        'Karipatti', 'Masinaickenpatti', 'Valasaiyur', 'Achankuttapatti', 'Kootathupatti', 'Arunoothumalai',
        'Thammampatti', 'Sendarapatti', 'Thedavur', 'Pachamalai', 'Goodamalai',
        'Panamarathupatti', 'Thumbalpatti', 'Mallur', 'Kondalampatti',
        'Ariyapalayam', 'Yethapur', 'Thumbal', 'Karumandurai', 'Soolankurichi', 'Kunnur',
        'Thalaivasal', 'Siruvachur', 'Veeraganur', 'Sathapadi', 'Kattukottai', 'Muttal',
        'Belur', 'Thirumanur', 'Valavanthi', 'Manjakuttai', 'Nagalur']
AREA = ['VP', 'TP', 'Mpty']
SEX = ['M', 'F', 'Mch', 'Fch']
CONDITIONS = ['Fever', 'Dengue NS1 Positive', 'Leptospirosis Positive', 'H1N1 Positive', 'Scrub Typhus Positive',
              'Typhoid Positive', 'Malaria Positive', 'Parainfluenza Positive']
TESTS = ['Dengue IgM Elisa Positive', 'Dengue NS1 Elisa Positive', 'Dengue NS1 & IgM Elisa Positive', 'Dengue Rapid NS1 Positive']
SOURCE = ['SSH', 'OVF']
MONTHS = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August', 'September', 'October', 'November', 'December']

# Column spec: (heading, width, kind, note). Headings must match the app's standard format exactly.
FEVER = [
    ('S.No', 6, 'int', None),
    ('Date of Admission', 13, 'date', 'DD-MM-YYYY, e.g. 03-09-2026'),
    ('Name', 22, 'text', None),
    ('Age ()', 8, 'age', 'Years as a number. Infants: e.g. 8 months, 14 days'),
    ('Sex   ( M /F )', 8, 'list:SEX', 'M / F (Mch / Fch for a child)'),
    ('Address : Village / Town', 34, 'text', None),
    ('Name of the HUD', 10, 'list:HUD', None),
    ('Name of the Block', 18, 'list:BLOCKS', 'Choose from the list'),
    ('Name of the PHC', 18, 'list:PHCS', 'Choose from the list or type the PHC name'),
    ('Contact Number', 13, 'phone', '10-digit mobile number'),
    ('Disease condition', 20, 'list:CONDITIONS', '"Fever", or the positive lab result'),
    ('Name of the Hospital / Institution', 22, 'text', None),
    ('Mpty / TP / VP', 8, 'list:AREA', None),
    ('Name of the local body', 20, 'text', 'Village panchayat, town panchayat or municipality'),
    ('Name of the ward', 10, 'text', 'Municipality / TP ward, e.g. Ward 9'),
    ('Name of the Habitation / Street Name', 22, 'text', None),
    ('Name of the HSC', 16, 'text', None),
    ('Date of On-set of fever', 13, 'date', 'DD-MM-YYYY'),
    ('Date of reporting', 13, 'date', 'DD-MM-YYYY. Month and Week NO fill in from this date'),
    ('Month', 11, 'auto_month:S', None),
    ('Week NO', 8, 'auto_week:S', None),
    ('Remarks', 16, 'text', None),
    ('Lattitude', 12, 'lat', 'Decimal degrees, 5 or more decimals, e.g. 11.59782'),
    ('Longitude', 12, 'lon', 'Decimal degrees, 5 or more decimals, e.g. 78.60303'),
    ('Outcome', 12, 'text', None),
    ('Remarks', 16, 'text', None),
]
DENGUE = [
    ('S.No', 6, 'int', None),
    ('Date', 13, 'date', 'Date of diagnosis, DD-MM-YYYY. Week no and Month fill in from this date'),
    ('Reporting SSH/ District', 14, 'text', None),
    ('Name', 22, 'text', None),
    ('Age', 8, 'age', 'Years as a number. Infants: e.g. 8 months'),
    ('Sex (M/F)', 8, 'list:SEX', None),
    ('ADDRESS', 34, 'text', None),
    ('Mobile Number', 13, 'phone', '10-digit mobile number'),
    ('Name of the Village Panchayat /Town Panchayat/ Municipality / Corporation', 24, 'text', None),
    ('Block', 18, 'list:BLOCKS', 'Choose from the list'),
    ('PHC / Urban PHC', 18, 'list:PHCS', 'Choose from the list or type the PHC name'),
    ('HUD', 10, 'list:HUD', None),
    ('Type of Test', 24, 'list:TESTS', None),
    ('Place of Diagnosis', 20, 'text', None),
    ('Mpty/ TP/VP', 8, 'list:AREA', None),
    ('Hamlet Ward', 16, 'text', None),
    ('SSH/ OVF', 8, 'list:SOURCE', None),
    ('Week no', 8, 'auto_week:B', None),
    ('Month', 11, 'auto_month:B', None),
    ('Health Sub-Centre', 16, 'text', None),
    ('Remarks', 16, 'text', None),
    ('Lat', 12, 'lat', 'Decimal degrees, e.g. 11.59782'),
    ('Long', 12, 'lon', 'Decimal degrees, e.g. 78.60303'),
]

wb = Workbook()
lists = wb.active
lists.title = 'Lists'
named = {'BLOCKS': BLOCKS, 'PHCS': PHCS, 'AREA': AREA, 'SEX': SEX, 'CONDITIONS': CONDITIONS, 'TESTS': TESTS,
         'SOURCE': SOURCE, 'HUD': ['Attur'], 'MONTHS': MONTHS}
refs = {}
for i, (key, vals) in enumerate(named.items(), start=1):
    col = get_column_letter(i)
    lists.cell(row=1, column=i, value=key).font = Font(bold=True)
    for r, v in enumerate(vals, start=2):
        lists.cell(row=r, column=i, value=v)
    refs[key] = f"Lists!${col}$2:${col}${len(vals) + 1}"
lists.sheet_state = 'hidden'


def date_parts(ref):
    return f'DATE(VALUE(RIGHT({ref},4)),VALUE(MID({ref},4,2)),VALUE(LEFT({ref},2)))'


def build(ws, spec):
    ws.freeze_panes = 'D2'
    ws.row_dimensions[1].height = 48
    last = get_column_letter(len(spec))
    ws.auto_filter.ref = f'A1:{last}{ROWS + 1}'
    for i, (head, width, kind, note) in enumerate(spec, start=1):
        col = get_column_letter(i)
        c = ws.cell(row=1, column=i, value=head)
        c.font, c.fill, c.border = HEAD_FONT, HEAD_FILL, BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.column_dimensions[col].width = width
        rng = f'{col}2:{col}{ROWS + 1}'
        dv = None
        if kind == 'date':
            for r in range(2, ROWS + 2):
                ws[f'{col}{r}'].number_format = '@'   # keep typed DD-MM-YYYY as text: Excel cannot swap day and month
            f2 = f'{col}2'
            dv = DataValidation(type='custom', allow_blank=True, showErrorMessage=True, errorStyle='stop',
                                formula1=f'AND(LEN({f2})=10,MID({f2},3,1)="-",MID({f2},6,1)="-",ISNUMBER({date_parts(f2)}))',
                                errorTitle='Date format', error='Type the date as DD-MM-YYYY, for example 03-09-2026.')
        elif kind.startswith('list:'):
            key = kind.split(':')[1]
            strict = key in ('BLOCKS', 'AREA', 'SOURCE', 'HUD')
            dv = DataValidation(type='list', formula1=refs[key], allow_blank=True, showErrorMessage=True,
                                errorStyle='stop' if strict else 'warning',
                                errorTitle='Not in the list',
                                error='Choose a value from the list.' if strict else 'This value is not in the list. Keep it only if it is correct.')
        elif kind in ('lat', 'lon'):
            lo, hi = (11.2, 12.1) if kind == 'lat' else (77.9, 79.0)
            for r in range(2, ROWS + 2):
                ws[f'{col}{r}'].number_format = '0.000000'
            dv = DataValidation(type='decimal', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                                showErrorMessage=True, errorStyle='warning', errorTitle='Check the coordinate',
                                error=f'{"Latitude" if kind == "lat" else "Longitude"} for Attur HUD is between {lo} and {hi}. '
                                      'Check that latitude and longitude are not swapped.')
        elif kind == 'phone':
            for r in range(2, ROWS + 2):
                ws[f'{col}{r}'].number_format = '@'
            f2 = f'{col}2'
            dv = DataValidation(type='custom', allow_blank=True, showErrorMessage=True, errorStyle='warning',
                                formula1=f'LEN({f2})>=10', errorTitle='Mobile number', error='A mobile number has 10 digits.')
        elif kind == 'int':
            dv = DataValidation(type='whole', operator='greaterThan', formula1='0', allow_blank=True)
        elif kind.startswith('auto_'):
            src = kind.split(':')[1]
            for r in range(2, ROWS + 2):
                s = f'{src}{r}'
                if kind.startswith('auto_month'):
                    f = f'=IF({s}="","",IFERROR(TEXT({date_parts(s)},"mmmm"),TEXT({s},"mmmm")))'
                else:
                    f = f'=IF({s}="","",IFERROR(_xlfn.ISOWEEKNUM({date_parts(s)}),_xlfn.ISOWEEKNUM({s})))'
                cell = ws[f'{col}{r}']
                cell.value = f
                cell.fill = AUTO_FILL
                cell.font = Font(color=GREY)
            note = (note or '') + 'Fills in automatically from the date.'
        if dv is not None:
            if note:
                dv.showInputMessage = True
                dv.promptTitle = head[:32]
                dv.prompt = note[:250]
            dv.add(rng)
            ws.add_data_validation(dv)
        elif note:
            dv2 = DataValidation(type=None, showInputMessage=True, promptTitle=head[:32], prompt=note[:250])
            dv2.add(rng)
            ws.add_data_validation(dv2)
    for r in range(2, 60):
        ws.row_dimensions[r].height = 18


fever = wb.create_sheet('Fever')
build(fever, FEVER)
dengue = wb.create_sheet('DENGUE')
build(dengue, DENGUE)

# Instructions
ins = wb.create_sheet('Instructions', 0)
ins.sheet_view.showGridLines = False
ins.column_dimensions['A'].width = 3
ins.column_dimensions['B'].width = 30
ins.column_dimensions['C'].width = 95
ins['B2'] = 'Attur HUD – Fever & Dengue line list'
ins['B2'].font = Font(size=18, bold=True, color=INK)
ins['B3'] = 'Standard template for the Fever & Dengue Surveillance Map. Keep one row per case.'
ins['B3'].font = Font(size=11, color=GREY)
rows = [
    ('Sheets', 'Fever = IP Fever cases (all admissions, with any lab result). DENGUE = confirmed dengue cases. '
               'The two may also be sent as separate files; the map combines them.'),
    ('Headings', 'Do not rename, delete or re-order the headings in row 1. Extra columns on the right are ignored.'),
    ('Dates', 'Type every date as DD-MM-YYYY, e.g. 03-09-2026. The date columns are set as text so Excel cannot '
              'swap day and month. A wrong format is refused.'),
    ('Month / Week', 'Grey columns fill in automatically from the date (ISO week, Monday–Sunday). Do not type over them.'),
    ('Block', 'Choose from the list: Attur, Attur Mpty, Narasingapuram Mpty, Ayothiyapattinam, Gangavalli, '
              'Panamarathupatti, Pethanaickenpalayam, Thalaivasal, Valapadi, Yercaud.'),
    ('PHC', 'Choose from the list (Attur HUD PHCs and UPHCs). Type a different name only if the PHC is missing.'),
    ('Location', 'Latitude and longitude in decimal degrees with at least 5 decimals, taken at the house. '
                 'Latitude ≈ 11.3–12.0, longitude ≈ 78.0–78.9 in Attur HUD. Never swap them.'),
    ('Age / Sex', 'Age in years as a number; for infants write "8 months" or "14 days". Sex: M, F, Mch, Fch.'),
    ('Disease condition', 'Fever sheet: "Fever", or the positive result such as "Leptospirosis Positive", '
                          '"H1N1 Positive", "Dengue NS1 Positive".'),
    ('Type of Test', 'DENGUE sheet: e.g. "Dengue IgM Elisa Positive", "Dengue NS1 Elisa Positive".'),
    ('Loading', 'Open the map, click Load Excel and choose this file (or drag it onto the page). '
                'The file is read on that computer only.'),
]
r = 5
for k, v in rows:
    ins.cell(row=r, column=2, value=k).font = Font(bold=True, color=TEAL, size=11)
    c = ins.cell(row=r, column=3, value=v)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ins.cell(row=r, column=2).alignment = Alignment(vertical='top')
    ins.row_dimensions[r].height = 32
    r += 1
r += 1
ins.cell(row=r, column=2, value='Example Fever row (for reference only – do not copy into the Fever sheet)').font = Font(bold=True, color=INK)
r += 1
example = [('Date of Admission', '02-09-2026'), ('Name', 'Example Name'), ('Age ()', '7'), ('Sex   ( M /F )', 'Fch'),
           ('Name of the Block', 'Valapadi'), ('Name of the PHC', 'Belur'), ('Disease condition', 'Fever'),
           ('Mpty / TP / VP', 'VP'), ('Name of the local body', 'Kurichi'), ('Date of reporting', '03-09-2026'),
           ('Lattitude', '11.651200'), ('Longitude', '78.401100')]
for k, v in example:
    ins.cell(row=r, column=2, value=k).font = Font(color=GREY)
    ins.cell(row=r, column=3, value=v)
    r += 1

wb.active = 1
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
if APP_COPY: shutil.copyfile(OUT, APP_COPY)
print('saved', OUT, os.path.getsize(OUT))
